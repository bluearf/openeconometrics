"""Admission checks and disposable parser processes for local file sources.

Reader allocation, retained batch buffers and total parser RSS are distinct.
The worker is a resource guard, not a security sandbox. Owned scratch is inside
the console scratch root when available, so Stop reclaims it with the session.
"""

from __future__ import annotations

import ctypes
import json
import multiprocessing
import os
from pathlib import Path
import sys
import tempfile
import time
import threading
import zipfile
import weakref

import pandas as pd

from openecon.data import DataError

MAX_CELL_BYTES = 1024 * 1024
MAX_RECORD_BYTES = 4 * 1024 * 1024
MAX_FOOTER_BYTES = 2 * 1024 * 1024
MAX_ROW_GROUP_BYTES = 256 * 1024 * 1024
MAX_READER_EXTRA_BYTES = 512 * 1024 * 1024
MAX_READER_RSS_BYTES = 1536 * 1024 * 1024
MAX_READER_BATCH_BYTES = 64 * 1024 * 1024
READER_TIMEOUT = 120


def rss_bytes(pid):
    """Current RSS of an owned parser, including shared runtime pages."""
    if sys.platform == "darwin":
        library = ctypes.CDLL("/usr/lib/libproc.dylib")
        buffer = ctypes.create_string_buffer(96)
        size = library.proc_pidinfo(int(pid), 4, 0, buffer, len(buffer))
        return int.from_bytes(buffer.raw[8:16], sys.byteorder) if size == len(buffer) else None
    if sys.platform.startswith("linux"):
        try:
            return int(Path(f"/proc/{pid}/statm").read_text().split()[1]) * os.sysconf(
                "SC_PAGE_SIZE"
            )
        except (OSError, IndexError):
            return None
    if os.name == "nt":
        from ctypes import wintypes

        class Counters(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD),
                ("faults", wintypes.DWORD),
                *[
                    (name, ctypes.c_size_t)
                    for name in (
                        "peak",
                        "working",
                        "paged_peak",
                        "paged",
                        "nonpaged_peak",
                        "nonpaged",
                        "pagefile",
                        "pagefile_peak",
                    )
                ],
            ]

        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel.CloseHandle.restype = wintypes.BOOL
        handle = kernel.OpenProcess(0x1000 | 0x10, False, int(pid))
        if not handle:
            return None
        try:
            value = Counters()
            value.cb = ctypes.sizeof(value)
            query = ctypes.WinDLL("psapi", use_last_error=True).GetProcessMemoryInfo
            query.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
            query.restype = wintypes.BOOL
            if query(handle, ctypes.byref(value), value.cb):
                return int(value.working)
        finally:
            kernel.CloseHandle(handle)
    return None


def peak_rss_bytes():
    """Peak for this parser's address space, excluding pre-exec parent history."""
    if sys.platform.startswith("linux"):
        # getrusage().ru_maxrss survives execve on Linux. A spawned parser can
        # inherit a console's historical peak even though it owns fresh mappings.
        # VmHWM belongs to the new address space and still catches brief peaks
        # that the parent's current-RSS polling did not observe.
        try:
            for line in Path("/proc/self/status").read_text().splitlines():
                if line.startswith("VmHWM:"):
                    _, value, unit = line.split()
                    peak = int(value)
                    return peak * 1024 if unit == "kB" and peak > 0 else None
        except (OSError, ValueError):
            return None
        return None
    if os.name == "posix":
        import resource

        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        return int(peak if sys.platform == "darwin" else peak * 1024)
    return None


def preflight_csv(path, *, header_only=False):
    """Bound fields/records using fixed byte blocks before any CSV parser call.

    Supports pandas' default comma/double-quote dialect, including multiline
    quoted fields, CRLF and escaped double quotes. Counts physical data records.
    """
    row_bytes = field_bytes = records = 0
    quoted = after_quote = False
    field_start = True
    deadline = time.monotonic() + READER_TIMEOUT
    with Path(path).open("rb") as stream:
        if stream.read(3) != b"\xef\xbb\xbf":
            stream.seek(0)
        while block := stream.read(65536):
            if time.monotonic() > deadline:
                raise DataError("CSV admission exceeded its time budget.", "READER_TIMEOUT")
            for byte in block:
                row_bytes += 1
                field_bytes += 1
                if row_bytes > MAX_RECORD_BYTES or field_bytes > MAX_CELL_BYTES:
                    raise DataError(
                        "CSV cell/record exceeds the parser allocation budget (1/4 MiB).",
                        "READER_LIMIT",
                    )
                if quoted:
                    if byte == 34:
                        quoted, after_quote = False, True
                    continue
                if after_quote and byte == 34:
                    quoted, after_quote = True, False
                    continue
                after_quote = False
                if byte == 34 and field_start:
                    quoted = True
                elif byte == 44:
                    field_bytes, field_start = 0, True
                    continue
                elif byte == 10:
                    records += 1
                    row_bytes = field_bytes = 0
                    field_start = True
                    if header_only:
                        return records
                    continue
                elif byte != 13:
                    field_start = False
        if quoted:
            raise DataError("CSV contains an unfinished quoted field.", "INVALID_SOURCE")
        if row_bytes:
            records += 1
    return max(0, records - 1)


def preflight_parquet(path):
    """Admit the physical footer length before Arrow interprets metadata."""
    path = Path(path)
    if path.stat().st_size < 12:
        raise DataError("Invalid Parquet file.", "INVALID_SOURCE")
    with path.open("rb") as stream:
        stream.seek(-8, 2)
        trailer = stream.read(8)
    size = int.from_bytes(trailer[:4], "little")
    if trailer[4:] != b"PAR1" or size > path.stat().st_size - 12:
        raise DataError("Invalid Parquet footer.", "INVALID_SOURCE")
    if size > MAX_FOOTER_BYTES:
        raise DataError(
            "Parquet footer exceeds the bounded parser metadata budget.", "READER_LIMIT"
        )
    return size


def preflight_xlsx(path):
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 10000:
            raise DataError("XLSX archive has too many members.", "READER_LIMIT")
        for entry in archive.infolist():
            # openpyxl materializes shared strings and styles even in read-only
            # mode. Worksheet XML itself is streamed inside the guarded worker.
            if (
                not entry.filename.startswith("xl/worksheets/")
                and entry.file_size > 32 * 1024 * 1024
            ):
                raise DataError(
                    "XLSX shared metadata exceeds its 32 MiB expansion budget.", "READER_LIMIT"
                )


def _child(connection, directory, mode, path, columns, batch_rows, dense_limits):
    parent = multiprocessing.parent_process()

    def parent_guard():
        while parent is not None and parent.is_alive():
            time.sleep(0.05)
        if parent is not None:
            os._exit(70)  # Console Stop owns/reclaims the enclosing scratch root.

    threading.Thread(target=parent_guard, daemon=True, name="File parser parent guard").start()

    def report(record):
        if os.name == "posix":
            peak = peak_rss_bytes()
            if not peak:
                raise DataError("Parser peak RSS supervision is unavailable.", "READER_LIMIT")
            record["peak_rss_bytes"] = peak
        connection.send(record)

    try:
        # Native allocation limit supplements RSS supervision on Linux. macOS
        # does not reliably enforce RLIMIT_AS; physical admission and RSS guard
        # remain in effect there. These limits apply only to this parser process.
        if sys.platform.startswith("linux"):
            import resource

            virtual = int(Path("/proc/self/statm").read_text().split()[0]) * os.sysconf(
                "SC_PAGE_SIZE"
            )
            resource.setrlimit(
                resource.RLIMIT_AS,
                (virtual + MAX_READER_EXTRA_BYTES, virtual + MAX_READER_EXTRA_BYTES),
            )
        report({"type": "ready"})
        if connection.recv() != "start":
            return
        if mode == "convert":
            record = _convert(Path(path), Path(directory) / "converted.parquet", batch_rows)
            report({"type": "converted", "record": record})
            return
        if mode == "read":
            import openecon.data as data_module

            data_module.MAX_FILE_BYTES, data_module.MAX_ROWS, data_module.MAX_COLUMNS = dense_limits
            from openecon.data import _read_local
            from openecon.dataset import Dataset

            try:
                value = _read_local(path)
            except DataError as exc:
                if exc.code != "DATA_LIMIT" or Path(path).suffix.lower() not in {".xlsx", ".dta"}:
                    raise
                record = _convert(Path(path), Path(directory) / "converted.parquet", batch_rows)
                report({"type": "converted", "record": record})
                return
            if isinstance(value, Dataset):
                report({"type": "dataset"})
            else:
                _write_frame(value, Path(directory) / "batch.arrow")
                report({"type": "frame"})
            return
        from openecon.dataset import scan

        source = scan(path)
        iterator = source._local_batches(tuple(columns), batch_rows)
        try:
            for frame in iterator:
                size = int(frame.memory_usage(index=True, deep=True).sum())
                if size > MAX_READER_BATCH_BYTES:
                    raise DataError(
                        "Parsed batch exceeds the separate 64 MiB parser buffer budget.",
                        "READER_LIMIT",
                    )
                target = Path(directory) / "batch.arrow"
                _write_frame(frame, target)
                report({"type": "batch", "bytes": target.stat().st_size})
                if connection.recv() != "next":
                    return
                target.unlink()
            report({"type": "done"})
        finally:
            iterator.close()
    except BaseException as exc:
        try:
            connection.send(
                {
                    "type": "error",
                    "code": getattr(exc, "code", "INVALID_SOURCE"),
                    "message": str(exc)[:2000],
                }
            )
        except (OSError, EOFError):
            pass
    finally:
        connection.close()


class Parser:
    def __init__(
        self, path, *, mode="batches", columns=(), batch_rows=8192, timeout=READER_TIMEOUT
    ):
        self.path, self.mode, self.columns, self.batch_rows = (
            str(path),
            mode,
            list(columns),
            batch_rows,
        )
        self.timeout, self.process, self.connection = timeout, None, None
        self.baseline = self.peak = 0
        scratch = os.environ.get("OPENECON_SCRATCH_DIRECTORY")
        self.temporary = tempfile.TemporaryDirectory(prefix=".openecon-reader-", dir=scratch)
        self.directory = Path(self.temporary.name)

    def start(self):
        parent, child = multiprocessing.get_context("spawn").Pipe()
        self.connection = parent
        self.process = multiprocessing.get_context("spawn").Process(
            target=_child,
            args=(
                child,
                str(self.directory),
                self.mode,
                self.path,
                self.columns,
                self.batch_rows,
                self._dense_limits(),
            ),
            name="OpenEconometrics file parser",
        )
        self.process.start()
        child.close()
        self.deadline = time.monotonic() + self.timeout
        if self.receive().get("type") != "ready":
            raise DataError("Parser did not start correctly.", "READER_LIMIT")
        self.baseline = rss_bytes(self.process.pid) or 0
        if not self.baseline:
            raise DataError(
                "Parser RSS supervision is unavailable on this platform.", "READER_LIMIT"
            )
        self.connection.send("start")
        return self

    @staticmethod
    def _dense_limits():
        from openecon.data import MAX_FILE_BYTES, MAX_ROWS, MAX_COLUMNS

        return MAX_FILE_BYTES, MAX_ROWS, MAX_COLUMNS

    def receive(self):
        while True:
            current = rss_bytes(self.process.pid)
            if current:
                self.peak = max(self.peak, current)
                if current > MAX_READER_RSS_BYTES or (
                    self.baseline and current - self.baseline > MAX_READER_EXTRA_BYTES
                ):
                    raise DataError(
                        "Parser exceeded its 512 MiB allocation / 1536 MiB process RSS budget.",
                        "READER_LIMIT",
                    )
            if time.monotonic() > self.deadline:
                raise DataError("Parser exceeded its time budget.", "READER_TIMEOUT")
            if self.connection.poll(0.01):
                try:
                    record = self.connection.recv()
                except EOFError as exc:
                    raise DataError(
                        "Parser ended before completing the source.", "READER_LIMIT"
                    ) from exc
                if record.get("type") == "error":
                    raise DataError(record["message"], record["code"])
                self.peak = max(self.peak, record.get("peak_rss_bytes", 0))
                if self.peak > MAX_READER_RSS_BYTES or (
                    self.baseline and self.peak - self.baseline > MAX_READER_EXTRA_BYTES
                ):
                    raise DataError("Parser peak RSS exceeded its resource budget.", "READER_LIMIT")
                return record
            if not self.process.is_alive():
                raise DataError("Parser stopped before returning a batch.", "READER_LIMIT")

    def close(self, *, cleanup=True):
        if self.connection:
            self.connection.close()
        if self.process:
            if self.process.is_alive():
                self.process.terminate()
            self.process.join(timeout=5)
            if self.process.is_alive():
                self.process.kill()
                self.process.join(timeout=5)
        if cleanup:
            self.temporary.cleanup()


def _write_frame(frame, target):
    import pyarrow as pa

    if int(frame.memory_usage(index=True, deep=True).sum()) > MAX_READER_BATCH_BYTES:
        raise DataError("Parser batch exceeds 64 MiB.", "READER_LIMIT")
    table = pa.Table.from_pandas(frame, preserve_index=None)
    attrs = json.dumps(frame.attrs, allow_nan=False, default=str).encode()
    if len(attrs) > 2 * 1024 * 1024:
        raise DataError("Batch metadata exceeds 2 MiB.", "METADATA_LIMIT")
    table = table.replace_schema_metadata(
        {**(table.schema.metadata or {}), b"openecon_batch_attrs": attrs}
    )
    with target.open("wb") as file:
        with pa.ipc.new_file(file, table.schema) as writer:
            writer.write_table(table)
    if target.stat().st_size > MAX_READER_BATCH_BYTES:
        raise DataError("Parser transfer exceeds 64 MiB.", "READER_LIMIT")


def _read_frame(path):
    import pyarrow as pa

    if path.stat().st_size > MAX_READER_BATCH_BYTES:
        raise DataError("Parser transfer exceeds 64 MiB.", "READER_LIMIT")
    with pa.ipc.open_file(path) as reader:
        table = reader.read_all()
        frame = table.to_pandas()
        frame.attrs = json.loads((table.schema.metadata or {}).get(b"openecon_batch_attrs", b"{}"))
    return frame


def read_source(path):
    from openecon.dataset import scan, _file_identity
    from openecon.frame import as_frame

    path = Path(path).resolve(strict=True)
    identity = _file_identity(path)
    if path.suffix.lower() == ".csv":
        preflight_csv(path)
    elif path.suffix.lower() == ".parquet":
        preflight_parquet(path)
    elif path.suffix.lower() == ".xlsx":
        preflight_xlsx(path)
    parser = Parser(path, mode="read")
    retain = False
    try:
        parser.start()
        record = parser.receive()
        if _file_identity(path) != identity:
            raise DataError("Source changed during parsing.", "SOURCE_CHANGED")
        if record["type"] == "frame":
            return as_frame(_read_frame(parser.directory / "batch.arrow"))
        if record["type"] == "dataset":
            return scan(path)
        if record["type"] != "converted":
            raise DataError("Invalid source parser result.", "READER_LIMIT")
        source = scan(parser.directory / "converted.parquet")
        if source.row_count != record["record"]["rows"]:
            raise DataError(
                "Converted physical row count differs from its receipt.", "SOURCE_CHANGED"
            )
        source._owned_reader_scratch = parser.temporary
        source._cleanup = weakref.finalize(source, parser.temporary.cleanup)
        source._provenance["conversion"] = {
            **record["record"],
            "original_path": str(path),
            "parser_peak_rss_bytes": parser.peak,
            "parser_baseline_rss_bytes": parser.baseline,
            "allocation_limit_bytes": MAX_READER_EXTRA_BYTES,
            "process_rss_limit_bytes": MAX_READER_RSS_BYTES,
            "batch_buffer_limit_bytes": MAX_READER_BATCH_BYTES,
        }
        retain = True
        return source
    finally:
        parser.close(cleanup=not retain)


def _convert(path, target, batch_rows):
    """Two bounded passes determine schema/count then emit physical Parquet groups."""
    from datetime import date, datetime
    import hashlib
    import pyarrow as pa
    import pyarrow.parquet as pq
    from openecon.dataset import _columns, _file_identity
    from openecon.data import _json_value

    identity = _file_identity(path)
    suffix = path.suffix.lower()
    batch_rows = min(8192, batch_rows)
    metadata, hidden = {}, {}
    if suffix == ".xlsx":
        preflight_xlsx(path)
        from openpyxl import load_workbook

        def xlsx_rows():
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                iterator = sheet.iter_rows(values_only=True)
                names = _columns(
                    [str(value) if value is not None else "" for value in next(iterator, ())]
                )
                yield names
                for row in iterator:
                    if len(row) > len(names) and any(
                        value is not None for value in row[len(names) :]
                    ):
                        raise DataError(
                            "XLSX row has values outside its declared columns.", "SCHEMA_MISMATCH"
                        )
                    row = tuple(row[: len(names)]) + (None,) * max(0, len(names) - len(row))
                    if any(
                        isinstance(value, str) and len(value.encode()) > MAX_CELL_BYTES
                        for value in row
                    ):
                        raise DataError("XLSX cell exceeds 1 MiB.", "READER_LIMIT")
                    yield row
            finally:
                workbook.close()

        iterator = xlsx_rows()
        columns = next(iterator)
        families = [set() for _ in columns]
        rows = 0
        try:
            for row in iterator:
                rows += 1
                for index, value in enumerate(row):
                    if value is None:
                        continue
                    family = (
                        "bool"
                        if isinstance(value, bool)
                        else "int"
                        if isinstance(value, int)
                        else "float"
                        if isinstance(value, float)
                        else "datetime"
                        if isinstance(value, datetime)
                        else "date"
                        if isinstance(value, date)
                        else "string"
                        if isinstance(value, str)
                        else "unsupported"
                    )
                    families[index].add(family)
        finally:
            iterator.close()
        types = []
        for family in families:
            if not family or family <= {"int"}:
                types.append(pa.int64())
            elif family <= {"int", "float"}:
                types.append(pa.float64())
            elif family == {"bool"}:
                types.append(pa.bool_())
            elif family <= {"date", "datetime"}:
                types.append(pa.timestamp("us"))
            elif family == {"string"}:
                types.append(pa.string())
            else:
                raise DataError(
                    "A mixed-type XLSX column has no lossless Parquet dtype; normalize it explicitly.",
                    "SCHEMA_MISMATCH",
                )
        schema = pa.schema(
            [pa.field(name, dtype) for name, dtype in zip(columns, types, strict=True)]
        )

        def frames():
            iterator = xlsx_rows()
            if next(iterator) != columns:
                raise DataError("XLSX columns changed between passes.", "SOURCE_CHANGED")
            block = []
            try:
                for row in iterator:
                    block.append(row)
                    if len(block) == batch_rows:
                        yield pa.Table.from_arrays(
                            [
                                pa.array([row[i] for row in block], type=types[i])
                                for i in range(len(columns))
                            ],
                            schema=schema,
                        )
                        block = []
                if block:
                    yield pa.Table.from_arrays(
                        [
                            pa.array([row[i] for row in block], type=types[i])
                            for i in range(len(columns))
                        ],
                        schema=schema,
                    )
            finally:
                iterator.close()
    elif suffix == ".dta":
        import pyreadstat

        _, meta = pyreadstat.read_dta(str(path), metadataonly=True, user_missing=True)
        columns = _columns(meta.column_names)
        if type(meta.number_rows) is not int or meta.number_rows < 0:
            raise DataError("DTA physical row count is unavailable.", "INVALID_SOURCE")
        rows = meta.number_rows
        metadata = _json_value(
            {
                "column_labels": meta.column_names_to_labels,
                "value_labels": meta.variable_value_labels,
                "original_types": meta.original_variable_types,
                "readstat_types": meta.readstat_variable_types,
            }
        )
        types = []
        for index, name in enumerate(columns):
            kind = meta.readstat_variable_types[name]
            display = meta.original_variable_types.get(name, "")
            if kind == "string":
                types.append(pa.string())
            else:
                types.append(
                    pa.date32()
                    if display.startswith("%td")
                    else pa.timestamp("us")
                    if display.startswith(("%tc", "%tC"))
                    else pa.int64()
                    if kind.startswith("int")
                    else pa.float64()
                )
                physical = f"__openecon_missing_{index}"
                while physical in columns:
                    physical += "_"
                hidden[name] = physical
        metadata["extended_missing_storage"] = hidden
        metadata["extended_missing_encoding"] = (
            "nullable code strings a-z at original row positions"
        )
        schema = pa.schema(
            [pa.field(name, dtype) for name, dtype in zip(columns, types, strict=True)]
            + [pa.field(name, pa.string()) for name in hidden.values()]
        )

        def frames():
            offset = 0
            while offset < rows:
                frame, current = pyreadstat.read_dta(
                    str(path),
                    user_missing=True,
                    row_offset=offset,
                    row_limit=min(batch_rows, rows - offset),
                )
                if tuple(frame.columns) != columns or len(frame) != min(batch_rows, rows - offset):
                    raise DataError(
                        "DTA physical rows/columns changed during conversion.", "SOURCE_CHANGED"
                    )
                arrays, codes = [], []
                for index, name in enumerate(columns):
                    series = frame[name]
                    if name in hidden:
                        allowed = current.missing_user_values.get(name, [])
                        mask = series.isin(allowed)
                        stored = series.where(mask, None).astype("string")
                        codes.append(pa.array(stored, type=pa.string(), from_pandas=True))
                        if pa.types.is_date(types[index]) or pa.types.is_timestamp(types[index]):
                            series = pd.to_datetime(series.mask(mask))
                        else:
                            series = pd.to_numeric(series.mask(mask), errors="raise")
                    arrays.append(pa.array(series, type=types[index], from_pandas=True))
                yield pa.Table.from_arrays([*arrays, *codes], schema=schema)
                offset += len(frame)
    else:
        raise DataError("Bounded conversion supports XLSX/DTA.", "UNSUPPORTED_FORMAT")
    metadata = _json_value(metadata)
    encoded = json.dumps({"metadata": metadata}, allow_nan=False).encode()
    if len(encoded) > 2 * 1024 * 1024:
        raise DataError("Converted labels/metadata exceed 2 MiB.", "METADATA_LIMIT")
    schema = schema.with_metadata({b"PANDAS_ATTRS": encoded})
    written = 0
    with pq.ParquetWriter(target, schema, compression="zstd") as writer:
        for table in frames():
            if table.nbytes > MAX_READER_BATCH_BYTES:
                raise DataError("Conversion block exceeds 64 MiB.", "READER_LIMIT")
            writer.write_table(
                table.replace_schema_metadata(schema.metadata), row_group_size=batch_rows
            )
            written += table.num_rows
            if target.stat().st_size > 8 * 1024 * 1024 * 1024:
                raise DataError("Owned conversion scratch exceeds 8 GiB.", "DISK_LIMIT")
    if written != rows or _file_identity(path) != identity:
        raise DataError("Source identity/row count changed during conversion.", "SOURCE_CHANGED")
    with pq.ParquetFile(target) as converted:
        if converted.metadata.num_rows != rows:
            raise DataError("Converted footer row count differs.", "SOURCE_CHANGED")
    with path.open("rb") as file:
        source_hash = hashlib.file_digest(file, "sha256").hexdigest()
    return dict(
        format=suffix.lstrip("."),
        rows=rows,
        columns=list(columns),
        batch_rows=batch_rows,
        source_sha256=source_hash,
        physical_output_bytes=target.stat().st_size,
        metadata=metadata,
        source_unchanged=True,
        converted_footer_verified=True,
    )


def guarded_batches(source, selected, batch_rows):
    if source._kind == "csv":
        preflight_csv(source._path)
    parser = Parser(source._path, columns=selected, batch_rows=min(batch_rows, 8192))
    try:
        parser.start()
        while True:
            record = parser.receive()
            if record["type"] == "done":
                break
            if record["type"] != "batch" or not 1 <= record["bytes"] <= MAX_READER_BATCH_BYTES:
                raise DataError("Invalid parser transfer.", "READER_LIMIT")
            path = parser.directory / "batch.arrow"
            if path.stat().st_size != record["bytes"]:
                raise DataError("Parser transfer changed.", "SOURCE_CHANGED")
            frame = _read_frame(path)
            source._reader_receipt = dict(
                parser_peak_rss_bytes=parser.peak,
                parser_baseline_rss_bytes=parser.baseline,
                reader_allocation_limit_bytes=MAX_READER_EXTRA_BYTES,
                parser_process_rss_limit_bytes=MAX_READER_RSS_BYTES,
                parser_buffer_limit_bytes=MAX_READER_BATCH_BYTES,
            )
            yield frame
            parser.deadline = time.monotonic() + parser.timeout
            parser.connection.send("next")
    finally:
        parser.close()
        source._reader_receipt = {
            **getattr(source, "_reader_receipt", {}),
            "parser_stopped": parser.process is None or not parser.process.is_alive(),
            "scratch_removed": not parser.directory.exists(),
        }
