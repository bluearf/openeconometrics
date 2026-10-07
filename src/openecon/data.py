"""Automatic bounded local loading and transparent dataset profiles."""
from __future__ import annotations

from datetime import date, datetime, timezone
import hashlib
import json
import math
import sys
from pathlib import Path
from uuid import uuid4

import pandas as pd

from openecon.frame import DataFrame, as_frame

MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100
SUPPORTED = {".csv", ".parquet", ".xlsx", ".dta"}
HASH_VERSION = 2
STREAM_HASH_VERSION = 3


class DataError(ValueError):
    def __init__(self, message: str, code: str = "INVALID_DATA"):
        super().__init__(message)
        self.code = code


def _json_value(value):
    """Normalize metadata as Parquet/JSON do, without dropping numeric precision."""
    # pandas can expose NumPy scalar values without OpenEconometrics importing NumPy.
    # Their public item() conversion preserves the existing canonical hash.
    if (type(value).__module__.split(".", 1)[0] == "numpy"
            and pd.api.types.is_scalar(value) and hasattr(value, "item")):
        value = value.item()
    if isinstance(value, dict):
        normalized = {str(key): _json_value(item) for key, item in value.items()}
        if len(normalized) != len(value):
            raise DataError("Metadata contains keys that collide when serialized.")
        return normalized
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if isinstance(value, (datetime, date)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    raise DataError("Dataset metadata and category labels must be JSON-compatible finite values.")


def _canonical_json(value) -> bytes:
    return json.dumps(_json_value(value), sort_keys=True, ensure_ascii=False,
                      separators=(",", ":"), allow_nan=False).encode("utf-8")


def metadata_hash(metadata: dict) -> str:
    return hashlib.sha256(_canonical_json(metadata)).hexdigest()


def frame_hash(frame: pd.DataFrame) -> str:
    """Bind values, column types, category semantics and metadata in row order.

    Row hashing is pandas-version-dependent. Profiles and replay scripts record
    that version and require it when checking this digest.
    """
    schema = []
    for name, dtype in frame.dtypes.items():
        column = {"name": str(name), "dtype": str(dtype)}
        if isinstance(dtype, pd.CategoricalDtype):
            column["categories"] = dtype.categories.tolist()
            column["ordered"] = dtype.ordered
        schema.append(column)
    digest = hashlib.sha256()
    digest.update(_canonical_json({"hash_version": HASH_VERSION, "schema": schema,
                                   "metadata": frame.attrs.get("metadata", {})}))
    digest.update(pd.util.hash_pandas_object(frame, index=True).to_numpy().astype("<u8", copy=False).tobytes())
    return digest.hexdigest()


def validate_frame(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        raise DataError("The dataset must contain at least one row and column.")
    if not frame.columns.is_unique:
        raise DataError("Column names must be unique.")
    frame = frame.copy()
    frame.columns = [str(c) for c in frame.columns]
    if not frame.columns.is_unique or any(not c.strip() for c in frame.columns):
        raise DataError("Column names must be nonempty and unique strings.")
    if any(len(c) > 200 for c in frame.columns):
        raise DataError("Column names cannot exceed 200 characters.")
    frame.reset_index(drop=True, inplace=True)
    if "metadata" in frame.attrs:
        frame.attrs["metadata"] = _json_value(frame.attrs["metadata"])
    return frame


def read(path: str | Path):
    """Read local data through a supervised parser; large XLSX/DTA become Dataset.

    Converted Parquet and missing-code storage are owned by the returned Dataset.
    Use Dataset.close() when finished. The original source is never modified.
    """
    if isinstance(path, str) and "://" in path:
        raise DataError("read accepts local paths only.", "INVALID_SOURCE")
    path = Path(path).expanduser()
    if path.is_dir():
        from openecon.dataset import scan
        return scan(path)
    if path.suffix.lower() not in SUPPORTED:
        raise DataError("Supported formats: CSV, Parquet, XLSX and DTA.", "UNSUPPORTED_FORMAT")
    if not path.is_file():
        raise DataError("The data file does not exist.", "FILE_NOT_FOUND")
    from openecon.source_readers import read_source
    return read_source(path)


def _read_local(path: str | Path):
    """Read a small table, or return a replayable Dataset for large CSV/Parquet.

    DTA labels and extended missing codes are retained in DataFrame.attrs.
    Extended numeric missings become NaN for estimation, with row/code metadata.
    100,000 rows, 100 columns and 32 MiB are dense-loading thresholds, not
    CSV/Parquet limits. XLSX/DTA must first be converted for bounded processing.
    """
    if isinstance(path, str) and "://" in path:
        raise DataError("read accepts local paths only; network URLs are not supported.", "INVALID_SOURCE")
    path = Path(path).expanduser()
    if path.is_dir():
        from openecon.dataset import scan
        return scan(path)
    if path.suffix.lower() not in SUPPORTED:
        raise DataError("Supported formats: CSV, Parquet, XLSX and DTA.", "UNSUPPORTED_FORMAT")
    if not path.is_file():
        raise DataError("The data file does not exist.", "FILE_NOT_FOUND")
    suffix = path.suffix.lower()
    if path.stat().st_size > MAX_FILE_BYTES:
        if suffix in {".csv", ".parquet"}:
            from openecon.dataset import scan
            return scan(path)
        raise DataError("Convert large XLSX/DTA files to CSV or Parquet for bounded processing.", "DATA_LIMIT")
    from openecon.optional_dependencies import require_extra
    if suffix == ".xlsx":
        require_extra("files", "openpyxl")
    elif suffix == ".dta":
        require_extra("files", "pyreadstat")
    try:
        if suffix == ".csv":
            # Inspect the raw header before pandas can silently rename duplicates.
            import csv
            with path.open(encoding="utf-8-sig", newline="") as stream:
                header = next(csv.reader(stream), [])
            if len(header) != len(set(header)) or any(not h.strip() for h in header):
                raise DataError("CSV headers must be nonempty and unique.")
            if len(header) > MAX_COLUMNS:
                from openecon.dataset import scan
                return scan(path)
            # Probe in bounded pieces before a dense parse. Read the small file
            # again to retain pandas' whole-column dtype inference semantics.
            count, memory = 0, 0
            with pd.read_csv(path, encoding="utf-8-sig", chunksize=min(8192, MAX_ROWS + 1),
                             dtype_backend="numpy_nullable") as reader:
                for piece in reader:
                    count += len(piece)
                    memory += int(piece.memory_usage(index=False, deep=True).sum())
                    if count > MAX_ROWS or memory > MAX_FILE_BYTES:
                        from openecon.dataset import scan
                        return scan(path)
            frame = pd.read_csv(path, encoding="utf-8-sig", dtype_backend="numpy_nullable")
        elif suffix == ".parquet":
            import pyarrow.parquet as pq
            metadata = pq.ParquetFile(path).metadata
            from openecon.source_readers import MAX_ROW_GROUP_BYTES
            if any(metadata.row_group(i).total_byte_size > MAX_ROW_GROUP_BYTES for i in range(metadata.num_row_groups)):
                raise DataError("Parquet row group exceeds its 256 MiB parser budget.", "READER_LIMIT")
            if metadata.num_rows > MAX_ROWS or metadata.num_columns > MAX_COLUMNS:
                from openecon.dataset import scan
                return scan(path)
            frame = pd.read_parquet(path)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                header = next(workbook.worksheets[0].iter_rows(max_row=1, values_only=True), ())
                names = [str(value) if value is not None else "" for value in header]
                if len(names) != len(set(names)) or any(not value.strip() for value in names):
                    raise DataError("XLSX headers must be nonempty and unique.")
            finally:
                workbook.close()
            frame = pd.read_excel(path, nrows=MAX_ROWS + 1, engine="openpyxl")
        else:
            import pyreadstat
            frame, meta = pyreadstat.read_dta(str(path), user_missing=True, row_limit=MAX_ROWS + 1)
            missing_codes = {}
            for col, codes in meta.missing_user_values.items():
                missing_codes[col] = {}
                for code in codes:
                    mask = frame[col].eq(code)
                    missing_codes[col][str(code)] = [position for position, matched in enumerate(mask) if matched]
                    frame.loc[mask, col] = float("nan")
                frame[col] = pd.to_numeric(frame[col], errors="raise")
            frame.attrs["metadata"] = {
                "column_labels": meta.column_names_to_labels,
                "value_labels": meta.variable_value_labels,
                "original_types": meta.original_variable_types,
                "extended_missing_codes": missing_codes,
            }
        if suffix in {".xlsx", ".dta"} and len(frame) > MAX_ROWS:
            raise DataError("Convert large XLSX/DTA files to CSV or Parquet for bounded processing.", "DATA_LIMIT")
        return as_frame(validate_frame(frame))
    except DataError:
        raise
    except Exception as exc:
        raise DataError(f"Could not read the {suffix} file: {type(exc).__name__}.") from exc


def example_frame() -> DataFrame:
    """Frozen synthetic teaching fixture, not observations about real people.

    These 480 rows are the original seed-2026 sample. Shipping the values keeps
    examples and existing dataset hashes stable across RNG/library changes.
    """
    frame = pd.read_csv(Path(__file__).with_name("examples") / "wages.csv",
                        float_precision="round_trip")
    frame.attrs["metadata"] = {"synthetic": True, "seed": 2026,
                               "description": "Synthetic wage and education teaching dataset."}
    return as_frame(frame)


def _number(value) -> float | None:
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def _preview(frame: pd.DataFrame, *, rows=50, budget=256 * 1024):
    """Return a display preview whose values and intermediate frame stay bounded."""
    selected = list(frame.columns[:MAX_COLUMNS])
    preview = frame.loc[:, selected].head(rows)
    while len(preview) and int(preview.memory_usage(index=True, deep=True).sum()) > budget:
        preview = preview.head(len(preview) // 2)
    records = json.loads(preview.to_json(orient="records", date_format="iso"))
    for row in records:
        for name, value in row.items():
            if isinstance(value, int) and abs(value) > 2**53 - 1:
                row[name] = str(value)
    while records and len(json.dumps(records, ensure_ascii=False).encode()) > budget:
        records.pop()
    return records, selected


def profile_frame(frame: pd.DataFrame, name: str = "Dataset", source: str = "upload",
                  dataset_id: str | None = None) -> dict:
    columns = []
    labels = frame.attrs.get("metadata", {}).get("column_labels", {})
    for col in frame:
        series = frame[col]
        numeric = pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_complex_dtype(series)
        finite = series.replace([float("inf"), -float("inf")], float("nan")) if numeric else None
        columns.append({
            "name": col, "label": labels.get(col), "dtype": str(series.dtype), "numeric": numeric,
            "missing": int(series.isna().sum()), "unique": int(series.nunique(dropna=True)),
            "mean": _number(finite.mean()) if numeric else None,
            "std": _number(finite.std()) if numeric else None,
            "min": _number(finite.min()) if numeric else None,
            "max": _number(finite.max()) if numeric else None,
        })
    # pandas JSON normalizes timestamps, NaN and NumPy scalar values safely.
    preview, preview_columns = _preview(frame)
    # JavaScript cannot display integers beyond 2**53 - 1 exactly. Only the
    # display preview uses strings; the stored frame retains its integer dtype.
    for row in preview:
        for column_name, value in row.items():
            if isinstance(value, int) and abs(value) > 2**53 - 1:
                row[column_name] = str(value)
    return {
        "id": dataset_id or str(uuid4()), "name": name, "source": source,
        "row_count": len(frame), "column_count": len(frame.columns), "columns": columns,
        "preview": preview, "data_hash": frame_hash(frame),
        "preview_columns": preview_columns, "preview_limit_bytes": 256 * 1024,
        "preview_integer_encoding": "Integers outside JavaScript's exact range are decimal strings.",
        "hash_version": HASH_VERSION, "hash_pandas_version": pd.__version__,
        "hash_algorithm": "SHA-256 over schema, category levels/order, metadata and pandas row hashes",
        "size_bytes": int(frame.memory_usage(index=True, deep=True).sum()),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "metadata": _json_value(frame.attrs.get("metadata", {})),
    }


def dataset_hash(data) -> str:
    """Verify a frame or immutable single-file Dataset without collecting rows."""
    from openecon.dataset import Dataset
    return data.file_hash() if isinstance(data, Dataset) else frame_hash(data)


def profile_dataset(data, name: str = "Dataset", source: str = "upload",
                    dataset_id: str | None = None, *, file_digest: str | None = None) -> dict:
    """Exact counts/moments and a bounded preview for a file-backed Dataset.

    Cardinality is exact only while bounded distinct-value storage fits. Capped
    columns expose a lower bound rather than presenting a sample as exact.
    """
    from openecon.dataset import Dataset
    if not isinstance(data, Dataset):
        raise DataError("profile_dataset requires a replayable Dataset.")
    if len(data._files) != 1 or data._path is None or data._path.is_dir():
        raise DataError("profile_dataset requires a single file-backed Dataset.", "INVALID_SOURCE")
    dataset_name = name
    states = {name: {"n": 0, "mean": 0., "m2": 0., "missing": 0,
                    "min": None, "max": None, "dtype": None, "numeric": True,
                    "unique": set(), "unique_bytes": 0, "unique_lower_bound": 0}
              for name in data.columns}
    total, distinct_bytes, preview_bytes = 0, 0, 0
    preview, preview_columns = [], data.columns[:MAX_COLUMNS]
    rows_per_batch = min(65_536, max(1, (8 * 1024 * 1024) // (64 * len(states))))
    for batch in data.iter_batches(batch_rows=rows_per_batch):
        total += len(batch)
        for column, state in states.items():
            values = batch[column]
            dtype = str(values.dtype)
            state["dtype"] = dtype if state["dtype"] is None else state["dtype"] if state["dtype"] == dtype else "mixed"
            state["missing"] += int(values.isna().sum())
            numeric = pd.api.types.is_numeric_dtype(values) and not pd.api.types.is_complex_dtype(values)
            state["numeric"] &= numeric
            if numeric:
                numeric_values = pd.to_numeric(values, errors="raise").astype("float64")
                finite = numeric_values[numeric_values.notna() & numeric_values.ne(float("inf")) & numeric_values.ne(-float("inf"))]
                count = len(finite)
                if count:
                    mean = float(finite.mean())
                    centered = finite - mean
                    m2 = float(centered.mul(centered).sum())
                    combined = state["n"] + count
                    delta = mean - state["mean"]
                    state["m2"] += m2 + delta * delta * (state["n"] / combined) * count
                    state["mean"] += delta * (count / combined)
                    state["n"] = combined
                    low, high = float(finite.min()), float(finite.max())
                    state["min"] = low if state["min"] is None else min(state["min"], low)
                    state["max"] = high if state["max"] is None else max(state["max"], high)
            if state["unique"] is not None:
                # Distinct values from one bounded reader block, never a full
                # source-column dictionary or array.
                for value in values.dropna().unique():
                    if pd.api.types.is_number(value) and not isinstance(value, complex):
                        text = str(int(value)) if math.isfinite(float(value)) and value == int(value) else str(value)
                        key = ("number", text)
                    else:
                        key = (type(value).__name__, str(value))
                    if key in state["unique"]:
                        continue
                    cost = sys.getsizeof(key) + sum(sys.getsizeof(item) for item in key) + 64
                    if len(state["unique"]) >= 256 or distinct_bytes + cost > 2 * 1024 * 1024:
                        state["unique_lower_bound"] = len(state["unique"]) + 1
                        distinct_bytes -= state["unique_bytes"]
                        state["unique"] = None
                        break
                    state["unique"].add(key)
                    state["unique_bytes"] += cost
                    distinct_bytes += cost
        if len(preview) < 50:
            records, _ = _preview(batch, rows=50 - len(preview), budget=256 * 1024 - preview_bytes)
            for row in records:
                for column, value in row.items():
                    if isinstance(value, int) and abs(value) > 2**53 - 1:
                        row[column] = str(value)
                cost = len(json.dumps(row, ensure_ascii=False).encode())
                if preview_bytes + cost <= 256 * 1024:
                    preview.append(row)
                    preview_bytes += cost
    if total == 0:
        raise DataError("The dataset must contain at least one row and column.")
    labels = data.metadata.get("column_labels", {})
    columns = []
    for name, state in states.items():
        columns.append({"name": name, "label": labels.get(name), "dtype": state["dtype"],
            "numeric": state["numeric"], "missing": state["missing"],
            "unique": len(state["unique"]) if state["unique"] is not None else None,
            "unique_exact": state["unique"] is not None,
            "unique_lower_bound": len(state["unique"]) if state["unique"] is not None else state["unique_lower_bound"],
            "mean": _number(state["mean"]) if state["numeric"] and state["n"] else None,
            "std": _number(math.sqrt(state["m2"] / (state["n"] - 1))) if state["numeric"] and state["n"] > 1 else None,
            "min": state["min"] if state["numeric"] else None,
            "max": state["max"] if state["numeric"] else None})
    return {"id": dataset_id or str(uuid4()), "name": dataset_name,
        "source": source, "row_count": total, "column_count": len(columns),
        "columns": columns, "preview": preview, "preview_columns": preview_columns,
        "preview_limit_bytes": 256 * 1024, "data_hash": file_digest or data.file_hash(),
        "hash_version": STREAM_HASH_VERSION, "hash_pandas_version": pd.__version__,
        "hash_algorithm": "SHA-256 over immutable source file bytes, including physical metadata",
        "size_bytes": data._files[0].stat().st_size,
        "created_at": datetime.now(timezone.utc).isoformat(), "metadata": _json_value(data.metadata),
        "processing": {"mode": "streaming", "batch_rows": rows_per_batch,
            "unique_budget_bytes": 2 * 1024 * 1024, "unique_limit_per_column": 256,
            "summary": "Exact row/missing counts; float64 finite-value moments over all rows (unrepresentable summaries are null); capped cardinalities are lower bounds."}}
